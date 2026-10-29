"""Build a real Word report from local metadata and genuine screenshot files.

Usage: python scripts/create-report.py [--template]
Requires python-docx. No credentials or screenshots are committed.
Without --template, refuses missing build URL, account, or screenshot evidence.
"""
import json
import sys
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_BREAK

root = Path(__file__).resolve().parent.parent
settings = json.loads((root / 'submission.local.json').read_text(encoding='utf-8-sig'))
template = '--template' in sys.argv
manifest_path = root / 'screenshots' / 'manifest.json'
manifest = json.loads(manifest_path.read_text(encoding='utf-8-sig')) if manifest_path.exists() else []
if not template:
    for key in ('expoBuildUrl', 'testEmail', 'testPassword'):
        if not settings.get(key):
            raise SystemExit(f'Missing {key} in submission.local.json')
    if not settings['expoBuildUrl'].startswith('https://'):
        raise SystemExit('Expo build URL must begin with https://')
    if not manifest:
        raise SystemExit('Add genuine screenshot paths and captions to screenshots/manifest.json first')
    for item in manifest:
        if not (root / 'screenshots' / item['file']).is_file() or not item.get('caption'):
            raise SystemExit(f'Missing real image or caption: {item}')

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(11.69)
section.top_margin = section.bottom_margin = Inches(0.8)
section.left_margin = section.right_margin = Inches(0.8)
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)
style.paragraph_format.space_after = Pt(8)
fields = [
    ('Student Code', settings['studentCode']),
    ('Full Name', settings['fullName']),
    ('GitHub Repository URL', settings['githubUrl']),
    ('Expo Project Link', settings.get('expoBuildUrl') or '[Add completed EAS build URL]'),
    ('Firebase Project ID', settings['firebaseProjectId']),
    ('Test account email', settings.get('testEmail') or '[Add existing test account email]'),
    ('Test account password', settings.get('testPassword') or '[Add test account password]'),
]
for label, value in fields:
    doc.add_paragraph(f'{label}: {value}')

for item in manifest:
    doc.add_page_break()
    caption = doc.add_paragraph(item['caption'])
    caption.paragraph_format.keep_with_next = True
    image = root / 'screenshots' / item['file']
    # Fit within printable A4 area while keeping the source image proportions.
    from PIL import Image
    with Image.open(image) as im:
        width, height = im.size
    available_width, available_height = 6.4, 8.6
    scale = min(available_width / width, available_height / height)
    doc.add_picture(str(image), width=Inches(width * scale), height=Inches(height * scale))
    if item.get('note'):
        doc.add_paragraph(item['note'])

out = root / 'submission'
out.mkdir(exist_ok=True)
filename = settings['studentCode'] + ('-template.docx' if template else '.docx')
doc.save(out / filename)
print(out / filename)
